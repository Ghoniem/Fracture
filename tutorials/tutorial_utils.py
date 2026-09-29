"""Shared helpers for the VCEM tutorials.

Every tutorial follows the same run contract as a full VCEM simulation:

    input/<case>.xlsx  ->  notebook  ->  output/YYYYMMDD_HHMMSS_<label>_<git-hash>/
                                          ├── provenance.md
                                          ├── summary.csv
                                          ├── diagnostics.txt
                                          ├── *.csv            (results tables)
                                          └── plots/*.png

The provenance file uses the four-section layout of RadCluster:
(1) input data, (2) user selections, (3) solver configuration,
(4) run statistics.
"""

from __future__ import annotations

import csv
import hashlib
import os
import platform
import subprocess
import sys
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Optional, Tuple

TUTORIALS_DIR = Path(__file__).resolve().parent
REPO_ROOT = TUTORIALS_DIR.parent
VCEM_DIR = REPO_ROOT / "vcem"
CPP_PY_DIR = VCEM_DIR / "cpp" / "python"

WORKBOOK_COLUMNS = ("Parameter", "Symbol", "Value", "Units", "Note")


# ── Paths and engine ─────────────────────────────────────────────────────────

def setup_paths() -> Path:
    """Put VCEM (fracture_utils) and the C++ extension dir on sys.path."""
    for p in (VCEM_DIR, CPP_PY_DIR):
        if str(p) not in sys.path:
            sys.path.insert(0, str(p))
    return VCEM_DIR


def resolve_engine(requested: str) -> Tuple[str, str]:
    """Return (engine_used, note). Falls back to 'python' if bem_cpp is missing."""
    requested = str(requested).lower().strip()
    if requested != "cpp":
        return "python", "python engine requested"
    try:
        import bem_cpp  # noqa: F401
        return "cpp", f"bem_cpp {getattr(bem_cpp, '__version__', '?')} imported"
    except Exception as exc:  # extension not built for this interpreter
        return "python", f"bem_cpp unavailable ({type(exc).__name__}); fell back to python"


class quiet:
    """Context manager that hides solver log lines (stdout) when enabled."""

    def __init__(self, enabled: bool = True):
        self.enabled = bool(enabled)

    def __enter__(self):
        if self.enabled:
            import contextlib
            import io
            self._cm = contextlib.redirect_stdout(io.StringIO())
            self._cm.__enter__()
        return self

    def __exit__(self, *exc):
        if self.enabled:
            self._cm.__exit__(*exc)
        return False


# ── Workbook I/O ─────────────────────────────────────────────────────────────

def _coerce(v: Any) -> Any:
    if isinstance(v, str):
        s = v.strip()
        try:
            f = float(s)
            return int(f) if f.is_integer() and "." not in s and "e" not in s.lower() else f
        except ValueError:
            return s
    if isinstance(v, float) and v.is_integer():
        return v
    return v


def write_workbook(path: Path, sheets: Mapping[str, Iterable[Tuple]]) -> Path:
    """Write a case workbook: one sheet per table, rows (Parameter, Symbol, Value, Units, Note)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="305496")
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        ws.append(list(WORKBOOK_COLUMNS))
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
        for row in rows:
            ws.append(list(row))
        for col, width in zip("ABCDE", (34, 18, 16, 12, 60)):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A2"
    wb.save(path)
    return path


def read_workbook(path: Path) -> Dict[str, Dict[str, Dict[str, Any]]]:
    """Read a case workbook into {sheet: {symbol: {value, units, parameter, note}}}."""
    from openpyxl import load_workbook

    wb = load_workbook(Path(path), data_only=True, read_only=True)
    tables: Dict[str, Dict[str, Dict[str, Any]]] = {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(h).strip() if h is not None else "" for h in rows[0]]
        idx = {h: i for i, h in enumerate(header)}
        missing = [c for c in ("Symbol", "Value") if c not in idx]
        if missing:
            raise ValueError(f"sheet {ws.title!r} lacks columns {missing}")
        table: Dict[str, Dict[str, Any]] = {}
        for r in rows[1:]:
            sym = r[idx["Symbol"]]
            if sym is None or str(sym).strip() == "":
                continue
            sym = str(sym).strip()
            if sym in table:
                raise ValueError(f"duplicate symbol {sym!r} in sheet {ws.title!r}")
            table[sym] = dict(
                value=_coerce(r[idx["Value"]]),
                units=r[idx["Units"]] if "Units" in idx else "",
                parameter=r[idx["Parameter"]] if "Parameter" in idx else sym,
                note=r[idx["Note"]] if "Note" in idx else "",
            )
        tables[ws.title] = table
    wb.close()
    return tables


def flat_values(tables: Mapping[str, Mapping[str, Mapping[str, Any]]]) -> Dict[str, Any]:
    """Collapse {sheet: {symbol: row}} to {symbol: value}; symbols must be unique."""
    out: Dict[str, Any] = {}
    for sheet, table in tables.items():
        for sym, row in table.items():
            if sym in out:
                raise ValueError(f"symbol {sym!r} appears in more than one sheet")
            out[sym] = row["value"]
    return out


def apply_overrides(tables, overrides: Optional[Mapping[str, Any]]) -> Dict[str, Tuple[Any, Any]]:
    """Replace workbook values in place for this run only. Returns {symbol: (old, new)}."""
    applied: Dict[str, Tuple[Any, Any]] = {}
    for sym, new in (overrides or {}).items():
        hits = [t for t in tables.values() if sym in t]
        if not hits:
            raise KeyError(f"override {sym!r} is not a Symbol in the workbook")
        row = hits[0][sym]
        applied[sym] = (row["value"], new)
        row["value"] = new
    return applied


def parse_list(v: Any, typ=float) -> List:
    """Parse a comma-separated workbook cell ('0, 15, 30') into a list."""
    if isinstance(v, (int, float)):
        return [typ(v)]
    return [typ(x) for x in str(v).replace(";", ",").split(",") if x.strip()]


def repo_relative(path: Path) -> str:
    """Path relative to the repository root (portable across machines)."""
    try:
        return Path(path).resolve().relative_to(REPO_ROOT).as_posix()
    except ValueError:
        return str(path)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


# ── Run directory and provenance ─────────────────────────────────────────────

def git_info() -> Dict[str, Any]:
    def _git(*args):
        return subprocess.run(["git", *args], cwd=REPO_ROOT, capture_output=True,
                              text=True, timeout=20).stdout.strip()
    try:
        return dict(
            git_sha=_git("rev-parse", "--short", "HEAD") or "nogit",
            git_branch=_git("rev-parse", "--abbrev-ref", "HEAD"),
            git_dirty=bool(_git("status", "--porcelain", "--untracked-files=no")),
        )
    except Exception:
        return dict(git_sha="nogit", git_branch="<unknown>", git_dirty="<unknown>")


def make_run_dir(tutorial_dir: Path, label: str) -> Tuple[Path, str]:
    """Create output/YYYYMMDD_HHMMSS_<label>_<git-hash>/plots and return (dir, run_label)."""
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    run_label = f"{ts}_{label}_{git_info()['git_sha']}"
    run_dir = Path(tutorial_dir) / "output" / run_label
    (run_dir / "plots").mkdir(parents=True, exist_ok=True)
    return run_dir, run_label


def machine_info() -> Dict[str, Any]:
    import numpy
    import scipy
    info = {
        "hostname": platform.node(),
        "platform": platform.platform(),
        "processor": platform.processor() or platform.machine(),
        "python": sys.version.split()[0],
        "numpy": numpy.__version__,
        "scipy": scipy.__version__,
        "cpu_count": os.cpu_count(),
        "omp_num_threads": os.environ.get("OMP_NUM_THREADS", "<unset>"),
    }
    try:
        import psutil
        info["ram_total_GB"] = round(psutil.virtual_memory().total / 1024**3, 2)
    except Exception:
        info["ram_total_GB"] = "<psutil unavailable>"
    return info


def _fmt_value(v: Any) -> str:
    if callable(v):
        return "<callable>"
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, float) or (isinstance(v, int) and abs(v) >= 1_000_000):
        return f"{v:.6g}"
    if isinstance(v, (tuple, list)):
        return "[" + ", ".join(_fmt_value(x) for x in v) + "]"
    return str(v)


def _write_dict_block(f, title: str, d: Mapping[str, Any]) -> None:
    f.write(f"### {title}\n\n")
    if not d:
        f.write("_(empty)_\n\n")
        return
    for k in sorted(d.keys(), key=str):
        f.write(f"- {k}: {_fmt_value(d[k])}\n")
    f.write("\n")


def write_provenance(run_dir: Path, run_label: str, *, title: str, tables,
                     overrides: Mapping[str, Tuple[Any, Any]],
                     user_selections: Mapping[str, Any],
                     solver_config: Mapping[str, Any],
                     run_stats: Mapping[str, Any]) -> Path:
    """Write provenance.md in the RadCluster four-section layout."""
    path = Path(run_dir) / "provenance.md"
    with open(path, "w", encoding="utf-8") as f:
        f.write(f"# {title} run — {run_label}\n\n")

        f.write("## (1) Input Data\n\n")
        f.write("_All input tables from the Excel workbook with parameter "
                "overrides applied._\n\n")
        for sheet, table in tables.items():
            _write_dict_block(f, sheet, {k: row["value"] for k, row in table.items()})
        _write_dict_block(f, "Overrides (workbook -> run)",
                          {k: f"{_fmt_value(a)} -> {_fmt_value(b)}" for k, (a, b) in overrides.items()})

        f.write("## (2) User Selections\n\n")
        _write_dict_block(f, "Run configuration", user_selections)

        f.write("## (3) Solver Configuration\n\n")
        _write_dict_block(f, "Solver settings", solver_config)

        f.write("## (4) Run Statistics\n\n")
        stats = dict(run_stats)
        try:
            import fracture_utils
            stats["vcem_version"] = fracture_utils.__version__
        except Exception:
            stats["vcem_version"] = "<unknown>"
        stats.update(git_info())
        stats.update(machine_info())
        _write_dict_block(f, "Runtime and machine", stats)
    return path


def write_csv(path: Path, rows: List[Mapping[str, Any]]) -> Path:
    """Write a list of dicts as CSV (columns in first-seen order; missing cells left blank)."""
    path = Path(path)
    if not rows:
        path.write_text("", encoding="utf-8")
        return path
    fields = list(dict.fromkeys(k for r in rows for k in r))
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for r in rows:
            w.writerow({k: (f"{v:.10g}" if isinstance(v, float) else v) for k, v in r.items()})
    return path


def write_diagnostics(path: Path, run_label: str, checks: List[Tuple[str, float, float, str]]) -> bool:
    """Write diagnostics.txt. Each check is (name, value, tolerance, units); passes if value <= tol."""
    all_ok = True
    lines = [f"Diagnostics — {run_label}", ""]
    for name, value, tol, units in checks:
        ok = bool(value <= tol)
        all_ok &= ok
        lines.append(f"[{'PASS' if ok else 'FAIL'}] {name}: {value:.4g} {units} (tolerance {tol:.4g} {units})")
    lines += ["", f"Overall: {'PASS' if all_ok else 'FAIL'}"]
    Path(path).write_text("\n".join(lines) + "\n", encoding="utf-8")
    return all_ok
